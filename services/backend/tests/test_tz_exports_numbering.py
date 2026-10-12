"""Letters and sheets quote the TZ document itself, under its own numbers."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from backend.enums import TZAnalysisStatus
from backend.schemas.analysis import TZAnalysisItem
from backend.services.analysis.docx_export import (
    build_clarification_docx,
    build_clarification_preview,
)
from backend.utils.requirements_struct import (
    normalize_tz_requirements,
    quoted_number,
    renumber_reference,
    requirement_numbers,
    strip_leading_key,
    tz_references,
)
from backend.utils.xlsx_export import build_tz_analysis_workbook

STATUS_LABELS = {status: status.value for status in TZAnalysisStatus}


def word_tz() -> dict:
    """Keys are positions: the point the document numbers «4.» has key 2."""
    return normalize_tz_requirements(
        {
            "1": {
                "text": "Общие требования",
                "label": "1.",
                "number": "1",
                "children": {
                    "1.1": {
                        "text": "Гарантия не менее 24 месяцев:",
                        "label": "1.1.",
                        "number": "1.1",
                        "page": 1,
                        "children": {
                            "1.1.1": {
                                "text": "на насос;",
                                "label": "1.1.1",
                                "number": "1.1.1",
                                "page": 1,
                            }
                        },
                    },
                },
            },
            "2": {
                "text": "Сроки",
                "label": "4.",
                "number": "4",
                "page": 2,
                "children": {
                    "2.1": {
                        "text": "Поставка за 90 дней",
                        "label": "4.1.",
                        "number": "4.1",
                        "page": 2,
                    },
                    "2.2": {
                        "text": "Допуск 3.5 мм",
                        "label": "",
                        "number": "",
                        "page": 2,
                    },
                },
            },
            "3": {
                "text": "Приложение 1",
                "label": "",
                "number": "",
                "page": 3,
                "children": {
                    "3.1": {
                        "text": "Торцевое уплотнение — 2 шт.",
                        "label": "1.",
                        "number": "1 (Приложение 1)",
                        "page": 3,
                    }
                },
            },
        }
    )


def item(requirement: str, **fields) -> TZAnalysisItem:
    return TZAnalysisItem(
        requirement=requirement,
        explanation="нет данных",
        status=TZAnalysisStatus.MISSING,
        **fields,
    )


# As the model really answers for a TZ read from Word: it copies the numbered
# line it was given into ref_value (internal key included) and refers to the
# point by that internal key.


def delivery_item() -> TZAnalysisItem:
    return item(
        "2.1. Поставка за 90 дней",
        ref="2.1",
        ref_value="2.1. Поставка за 90 дней",
        requirement_ref="Источник ТЗ, стр 2, пункт 2.1",
    )


def warranty_item() -> TZAnalysisItem:
    return item(
        "1.1.1. на насос;\n  (исходный пункт: Гарантия не менее 24 месяцев:)",
        ref="1.1.1",
        ref_value="Гарантия не менее 24 месяцев: 1.1.1. на насос;",
        requirement_ref="Источник ТЗ, стр 1, пункт 1.1.1",
    )


def tolerance_item() -> TZAnalysisItem:
    return item(
        "2.2. Допуск 3.5 мм",
        ref="2.2",
        ref_value="2.2. Допуск 3.5 мм",
        requirement_ref="Источник ТЗ, стр 2, пункт 2.2",
    )


def appendix_item() -> TZAnalysisItem:
    return item(
        "3.1. Торцевое уплотнение — 2 шт.",
        ref="3.1",
        ref_value="3.1. Торцевое уплотнение — 2 шт.",
        requirement_ref="Источник ТЗ, стр 3, пункт 3.1",
    )


def letter(items: list[TZAnalysisItem], tz: dict | None) -> list[str]:
    _, paragraphs, _ = build_clarification_preview(
        items, list(range(len(items))), requirements_tz=tz
    )
    return paragraphs


def test_the_letter_quotes_the_point_under_the_number_of_the_document() -> (
    None
):
    paragraphs = letter([delivery_item()], word_tz())

    assert (
        'Требование: "4.1. Поставка за 90 дней" (Источник ТЗ, стр 2, пункт 4.1)'
        in paragraphs
    )


def test_a_point_is_led_by_the_line_that_introduces_it() -> None:
    paragraphs = letter([warranty_item()], word_tz())

    assert (
        'Требование: "1.1.1. Гарантия не менее 24 месяцев: на насос;" '
        "(Источник ТЗ, стр 1, пункт 1.1.1)"
    ) in paragraphs


def test_a_point_the_document_does_not_number_is_quoted_by_page() -> None:
    paragraphs = letter([tolerance_item()], word_tz())

    assert 'Требование: "Допуск 3.5 мм" (Источник ТЗ, стр 2)' in paragraphs


def test_a_number_the_document_repeats_says_where_it_belongs() -> None:
    paragraphs = letter([appendix_item()], word_tz())

    assert (
        'Требование: "1 (Приложение 1). Торцевое уплотнение — 2 шт." '
        "(Источник ТЗ, стр 3, пункт 1 (Приложение 1))"
    ) in paragraphs


def test_no_internal_key_reaches_the_letter() -> None:
    items = [
        delivery_item(),
        warranty_item(),
        tolerance_item(),
        appendix_item(),
    ]
    text = "\n".join(letter(items, word_tz()))

    # positions that differ from the number the document gives the point
    for internal in ("2.1.", "2.2.", "3.1.", "пункт 2.1", "пункт 3.1"):
        assert internal not in text, internal


def test_the_letter_is_unchanged_for_a_tz_found_by_the_model() -> None:
    legacy = normalize_tz_requirements(
        {"2": {"text": "Сроки", "children": {"2.1": {"text": "Поставка"}}}}
    )
    clause = item(
        "2.1. Поставка",
        ref="2",
        ref_value="Сроки",
        requirement_ref="Источник ТЗ, стр 2, пункт 2",
    )

    for tz in (None, legacy):
        paragraphs = letter([clause], tz)
        assert (
            'Требование: "2. Сроки" (Источник ТЗ, стр 2, пункт 2)'
            in paragraphs
        )


def test_the_word_letter_quotes_the_same_text() -> None:
    data = build_clarification_docx(
        [delivery_item()], [0], requirements_tz=word_tz()
    )
    lines = [p.text for p in Document(BytesIO(data)).paragraphs]

    assert (
        'Требование: "4.1. Поставка за 90 дней" (Источник ТЗ, стр 2, пункт 4.1)'
        in lines
    )


def test_the_sheet_quotes_the_same_text() -> None:
    wb = build_tz_analysis_workbook(
        [delivery_item(), tolerance_item(), appendix_item()],
        STATUS_LABELS,
        requirements_tz=word_tz(),
    )
    ws = wb.active
    assert ws is not None
    rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]

    assert [(row[2], row[3]) for row in rows] == [
        ("4.1. Поставка за 90 дней", "Источник ТЗ, стр 2, пункт 4.1"),
        ("Допуск 3.5 мм", "Источник ТЗ, стр 2"),
        (
            "1 (Приложение 1). Торцевое уплотнение — 2 шт.",
            "Источник ТЗ, стр 3, пункт 1 (Приложение 1)",
        ),
    ]


def test_number_helpers() -> None:
    numbers = requirement_numbers(word_tz())
    assert numbers == {
        "1": "1",
        "1.1": "1.1",
        "1.1.1": "1.1.1",
        "2": "4",
        "2.1": "4.1",
        "2.2": None,
        "3": None,
        "3.1": "1 (Приложение 1)",
    }
    assert quoted_number(numbers, "2") == "4"
    assert quoted_number(numbers, "2.2") is None
    assert quoted_number(numbers, "9.9") == "9.9"  # no longer in the TZ
    assert quoted_number(None, "2") == "2"
    assert requirement_numbers({"1": {"text": "x", "children": {}}}) is None

    assert renumber_reference("Источник ТЗ, стр 3, пункт 2", numbers) == (
        "Источник ТЗ, стр 3, пункт 4"
    )
    assert renumber_reference("Источник ТЗ, стр 3, пункт 2.2", numbers) == (
        "Источник ТЗ, стр 3"
    )
    assert renumber_reference("Источник ТЗ, пункт 2", None) == (
        "Источник ТЗ, пункт 2"
    )
    assert strip_leading_key("2.1. Допуск") == "Допуск"
    assert strip_leading_key("3.5 мм допуск") == "3.5 мм допуск"


def test_quotes_never_hold_an_internal_key() -> None:
    references = tz_references(word_tz())
    assert references is not None
    assert references.quotes == {
        "1.1.1": "Гарантия не менее 24 месяцев: на насос;",
        "2.1": "Поставка за 90 дней",
        "2.2": "Допуск 3.5 мм",
        "3.1": "Торцевое уплотнение — 2 шт.",
    }
    legacy = normalize_tz_requirements({"1": {"text": "x", "children": {}}})
    assert tz_references(legacy) is None
